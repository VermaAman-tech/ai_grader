from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from db.database import session_scope
from db.models import Course, Exam, Grade, Rubric, Student, Submission


def export_exam_gradebook(exam_id: int, output_path: str) -> str:
    with session_scope() as db:
        exam = db.get(Exam, exam_id)
        if not exam:
            raise ValueError(f"Exam {exam_id} not found")

        course = db.get(Course, exam.course_id)
        students = (
            db.query(Student)
            .filter(Student.course_id == exam.course_id)
            .order_by(Student.roll_number.asc().nullslast(), Student.name.asc())
            .all()
        )
        rubrics = (
            db.query(Rubric)
            .filter(Rubric.exam_id == exam_id)
            .order_by(Rubric.question_order.asc(), Rubric.id.asc())
            .all()
        )
        submissions = db.query(Submission).filter(Submission.exam_id == exam_id).all()
        grade_rows = (
            db.query(Grade)
            .join(Submission, Submission.id == Grade.submission_id)
            .filter(Submission.exam_id == exam_id)
            .all()
        )

        submission_by_student = {submission.student_id: submission for submission in submissions}
        grade_lookup = {(grade.submission_id, grade.rubric_id): grade for grade in grade_rows}

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Gradebook"

        headers = ["Roll Number", "Student Name"]
        for rubric in rubrics:
            headers.append(f"{rubric.question_no} (/{rubric.max_marks})")
        headers.extend(["Total", "Percentage", "Status"])

        header_fill = PatternFill("solid", fgColor="162638")
        header_font = Font(color="F3FBFF", bold=True)
        border = Border(
            left=Side(style="thin", color="43556C"),
            right=Side(style="thin", color="43556C"),
            top=Side(style="thin", color="43556C"),
            bottom=Side(style="thin", color="43556C"),
        )

        for col_idx, title in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=col_idx, value=title)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border

        max_total = sum(rubric.max_marks for rubric in rubrics)

        for row_idx, student in enumerate(students, start=2):
            submission = submission_by_student.get(student.id)
            row_values = [student.roll_number or "", student.name]
            total = 0.0

            for rubric in rubrics:
                mark_value = 0.0
                if submission:
                    grade = grade_lookup.get((submission.id, rubric.id))
                    if grade:
                        mark_value = grade.override_marks if grade.override_marks is not None else grade.awarded_marks
                row_values.append(round(mark_value, 2))
                total += mark_value

            percentage = (total / max_total * 100.0) if max_total > 0 else 0.0
            status = submission.status if submission else "missing"
            row_values.extend([round(total, 2), round(percentage, 2), status])

            for col_idx, value in enumerate(row_values, start=1):
                cell = sheet.cell(row=row_idx, column=col_idx, value=value)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = border

            pct_cell = sheet.cell(row=row_idx, column=len(headers) - 1)
            if percentage >= 75:
                pct_cell.fill = PatternFill("solid", fgColor="1D5C3A")
            elif percentage >= 40:
                pct_cell.fill = PatternFill("solid", fgColor="6A5313")
            else:
                pct_cell.fill = PatternFill("solid", fgColor="6B1F2A")

        for col_idx in range(1, len(headers) + 1):
            max_len = 12
            for row in range(1, sheet.max_row + 1):
                value = sheet.cell(row=row, column=col_idx).value
                max_len = max(max_len, len(str(value or "")) + 2)
            sheet.column_dimensions[get_column_letter(col_idx)].width = min(max_len, 38)

        sheet.freeze_panes = "A2"
        output = Path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output)

        details_sheet = workbook.create_sheet("Metadata")
        details_sheet["A1"] = "Course"
        details_sheet["B1"] = f"{course.code} - {course.name}" if course else ""
        details_sheet["A2"] = "Exam"
        details_sheet["B2"] = exam.name
        details_sheet["A3"] = "Total students"
        details_sheet["B3"] = len(students)
        details_sheet["A4"] = "Rubric questions"
        details_sheet["B4"] = len(rubrics)
        workbook.save(output)

    return str(output)
